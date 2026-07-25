from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import load_workbook


BRACS_LABELS = ("N", "PB", "UDH", "FEA", "ADH", "DCIS", "IC")
_BRACS_FILE_RE = re.compile(r"^(BRACS_\d+)\.(svs|qpdata)$", re.IGNORECASE)


@dataclass(slots=True)
class BracsCaseRecord:
    excel_row_index: int
    case_id: str
    gt_label: str
    split: str
    patient_id: str | None
    declared_roi_count: int | None
    svs_path: Path | None
    qpdata_path: Path | None
    duplicate_svs_paths: list[str] = field(default_factory=list)
    duplicate_qpdata_paths: list[str] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def is_runnable(self) -> bool:
        return self.svs_path is not None

    @property
    def has_qpdata(self) -> bool:
        return self.qpdata_path is not None

    def summary_dict(self) -> dict[str, Any]:
        return {
            "excel_row_index": self.excel_row_index,
            "case_id": self.case_id,
            "gt_label": self.gt_label,
            "split": self.split,
            "patient_id": self.patient_id,
            "declared_roi_count": self.declared_roi_count,
            "svs_path": str(self.svs_path) if self.svs_path else None,
            "qpdata_path": str(self.qpdata_path) if self.qpdata_path else None,
            "duplicate_svs_paths": list(self.duplicate_svs_paths),
            "duplicate_qpdata_paths": list(self.duplicate_qpdata_paths),
            "metadata": dict(self.metadata),
        }


@dataclass(slots=True)
class BracsDatasetManifest:
    data_root: Path
    excel_path: Path
    cases: list[BracsCaseRecord]
    indexed_svs_count: int
    indexed_qpdata_count: int
    duplicate_filename_counts: dict[str, int] = field(default_factory=dict)

    def runnable_cases(self, *, allowed_splits: Iterable[str] | None = None) -> list[BracsCaseRecord]:
        allowed = _normalize_split_filter(allowed_splits)
        cases = [case for case in self.cases if case.is_runnable]
        if allowed is not None:
            cases = [case for case in cases if case.split.strip().lower() in allowed]
        return cases

    def skipped_missing_svs_cases(self, *, allowed_splits: Iterable[str] | None = None) -> list[BracsCaseRecord]:
        allowed = _normalize_split_filter(allowed_splits)
        cases = [case for case in self.cases if not case.is_runnable]
        if allowed is not None:
            cases = [case for case in cases if case.split.strip().lower() in allowed]
        return cases

    def summary_dict(self) -> dict[str, Any]:
        return {
            "data_root": str(self.data_root),
            "excel_path": str(self.excel_path),
            "n_cases_in_excel": len(self.cases),
            "indexed_svs_count": self.indexed_svs_count,
            "indexed_qpdata_count": self.indexed_qpdata_count,
            "n_runnable_cases": sum(1 for case in self.cases if case.is_runnable),
            "n_cases_with_qpdata": sum(1 for case in self.cases if case.has_qpdata),
            "duplicate_filename_counts": dict(self.duplicate_filename_counts),
        }


def load_bracs_manifest(data_root: str | Path) -> BracsDatasetManifest:
    root = Path(data_root).expanduser().resolve()
    excel_path = root / "BRACS.xlsx"
    if not excel_path.exists():
        raise FileNotFoundError(f"BRACS.xlsx was not found at the dataset root: {excel_path}")

    indexed = _index_case_files(root)
    duplicate_filename_counts = {
        name: len(paths)
        for name, paths in {**indexed["svs"], **indexed["qpdata"]}.items()
        if len(paths) > 1
    }

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if "WSI_Information" not in workbook.sheetnames:
            raise ValueError("BRACS.xlsx is missing the WSI_Information sheet.")
        sheet = workbook["WSI_Information"]
        headers = [str(cell).strip() if cell is not None else "" for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        header_index = {name: idx for idx, name in enumerate(headers)}
        required = ["WSI Filename", "Patient Id", "RoI", "WSI label", "Set"]
        missing = [name for name in required if name not in header_index]
        if missing:
            raise ValueError(f"WSI_Information is missing required columns: {missing}")

        cases: list[BracsCaseRecord] = []
        for excel_row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            case_stem = _normalize_case_stem(row[header_index["WSI Filename"]])
            if not case_stem:
                continue
            gt_label = normalize_bracs_label(row[header_index["WSI label"]])
            if gt_label is None:
                raise ValueError(f"Unsupported BRACS label at Excel row {excel_row_index}: {row[header_index['WSI label']]}")
            split = str(row[header_index["Set"]] or "").strip()
            patient_id = _string_or_none(row[header_index["Patient Id"]])
            declared_roi_count = _int_or_none(row[header_index["RoI"]])

            svs_name = f"{case_stem}.svs".lower()
            qpdata_name = f"{case_stem}.qpdata".lower()
            svs_matches = indexed["svs"].get(svs_name, [])
            qpdata_matches = indexed["qpdata"].get(qpdata_name, [])

            cases.append(
                BracsCaseRecord(
                    excel_row_index=excel_row_index,
                    case_id=case_stem,
                    gt_label=gt_label,
                    split=split,
                    patient_id=patient_id,
                    declared_roi_count=declared_roi_count,
                    svs_path=svs_matches[0] if svs_matches else None,
                    qpdata_path=qpdata_matches[0] if qpdata_matches else None,
                    duplicate_svs_paths=[str(path) for path in svs_matches[1:]],
                    duplicate_qpdata_paths=[str(path) for path in qpdata_matches[1:]],
                    metadata={
                        "excel_row_index": excel_row_index,
                    },
                )
            )
    finally:
        workbook.close()

    return BracsDatasetManifest(
        data_root=root,
        excel_path=excel_path,
        cases=cases,
        indexed_svs_count=sum(len(paths) for paths in indexed["svs"].values()),
        indexed_qpdata_count=sum(len(paths) for paths in indexed["qpdata"].values()),
        duplicate_filename_counts=duplicate_filename_counts,
    )


def normalize_bracs_label(value: Any) -> str | None:
    text = _normalize_bracs_label_text(value)
    if not text:
        return None
    collapsed = text.replace("-", "")
    aliases = {
        "N": "N",
        "PB": "PB",
        "UDH": "UDH",
        "FEA": "FEA",
        "ADH": "ADH",
        "DCIS": "DCIS",
        "IC": "IC",
        "BENIGN": "N",
        "PATHOLOGICALBENIGN": "PB",
    }
    normalized = aliases.get(collapsed)
    if normalized is not None:
        return normalized

    tokenized = text.replace("-", " ")
    for label in sorted(BRACS_LABELS, key=len, reverse=True):
        if re.search(rf"\b{re.escape(label)}\b", tokenized):
            return label
    return None


def _normalize_bracs_label_text(value: Any) -> str:
    text = str(value or "").strip().upper()
    if not text:
        return ""
    text = re.sub(r"[^A-Z0-9]+", "-", text).strip("-")
    while text.endswith("-SURE"):
        text = text[: -len("-SURE")].rstrip("-")
    return text


def _normalize_case_stem(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    match = re.match(r"^(BRACS_\d+)(?:\.svs)?$", text, flags=re.IGNORECASE)
    if not match:
        return text
    return match.group(1).upper().replace("BRACS_", "BRACS_")


def _normalize_split_filter(allowed_splits: Iterable[str] | None) -> set[str] | None:
    if allowed_splits is None:
        return None
    normalized = {str(item).strip().lower() for item in allowed_splits if str(item).strip()}
    return normalized or None


def _index_case_files(root: Path) -> dict[str, dict[str, list[Path]]]:
    svs_index: dict[str, list[Path]] = {}
    qpdata_index: dict[str, list[Path]] = {}
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        match = _BRACS_FILE_RE.match(path.name)
        if not match:
            continue
        key = path.name.lower()
        suffix = match.group(2).lower()
        target = svs_index if suffix == "svs" else qpdata_index
        target.setdefault(key, []).append(path.resolve())
    return {"svs": svs_index, "qpdata": qpdata_index}


def _string_or_none(value: Any) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None


def _int_or_none(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except Exception:
        try:
            return int(float(value))
        except Exception:
            return None
